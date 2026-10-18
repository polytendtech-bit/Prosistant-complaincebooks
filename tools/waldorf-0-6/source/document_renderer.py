from pathlib import Path
import json,csv,datetime
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
R=Path(__file__).resolve().parents[1];d=json.loads((R/'curriculum.json').read_text());refs=json.loads((R/'source/RESEARCH_ATTEMPTS.json').read_text());A={a['id']:a for a in d['areas']};B={b['id']:b for b in d['ageBands']}
for name,file in [('ELS','DejaVuSans.ttf'),('ELB','DejaVuSans-Bold.ttf'),('ELT','DejaVuSerif.ttf')]:pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/dejavu/'+file))
ink=colors.HexColor('#243d35');green=colors.HexColor('#e7efe7');line=colors.HexColor('#bfcbbf')
styles={k:ParagraphStyle(k,fontName=font,fontSize=size,leading=leading,textColor=ink,spaceAfter=after) for k,font,size,leading,after in [('title','ELT',24,29,14),('body','ELS',9.4,14,9),('small','ELS',7.5,11,5),('heading','ELB',12,17,9),('table','ELS',8,11.5,0)]}
def P(s,k='body'):return Paragraph(escape(str(s)).replace('\n','<br/>'),styles[k])
def footer(c,doc):c.setFont('ELS',7);c.setFillColor(ink);c.drawString(42,24,'Waldorf-inspired curriculum · Birth through age six');c.drawRightString(A4[0]-42,24,str(doc.page))
def export_book(name,pages):
 story=[];doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.65);sec.right_margin=Inches(.65);normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(7)
 for i,page in enumerate(pages):
  if i:story.append(PageBreak());doc.add_page_break()
  story.append(P(page['title'],'title'));doc.add_heading(page['title'],0)
  for block in page['blocks']:
   kind=block[0]
   if kind in ['p','small','h']:
    text=block[1];story.append(P(text,{'p':'body','small':'small','h':'heading'}[kind]));doc.add_heading(text,2) if kind=='h' else doc.add_paragraph(text)
   elif kind=='table':
    headers,rows,widths=block[1:];data=[[P(x,'table') for x in headers]]+[[P(x,'table') for x in row] for row in rows];t=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.35,line),('BACKGROUND',(0,0),(-1,0),green),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]));story.append(t);story.append(Spacer(1,9));table=doc.add_table(rows=1,cols=len(headers));table.style='Table Grid'
    for cell,text in zip(table.rows[0].cells,headers):cell.text=text
    for row in rows:
     for cell,text in zip(table.add_row().cells,row):cell.text=str(text)
   elif kind=='lines':
    label,n=block[1:];story.append(P(label,'heading'));t=Table([['']]*n,colWidths=[500],rowHeights=22);t.setStyle(TableStyle([('LINEBELOW',(0,0),(-1,-1),.35,line)]));story.append(t);story.append(Spacer(1,8));doc.add_heading(label,2)
    for _ in range(n):doc.add_paragraph('_'*75)

 pdf=SimpleDocTemplate(str(R/(name+'.pdf')),pagesize=A4,leftMargin=42,rightMargin=42,topMargin=38,bottomMargin=42,title='Waldorf-inspired — '+name.replace('_',' '),author='Growing Together');pdf.build(story,onFirstPage=footer,onLaterPages=footer);doc.save(R/(name+'.docx'))
 return len(pages)
def page(title,*blocks):return {'title':title,'blocks':list(blocks)}
