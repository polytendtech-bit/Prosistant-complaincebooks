import json,csv,pathlib,collections,html,re,zipfile,hashlib
from docx import Document
from docx.shared import Pt, Inches
base=pathlib.Path('/workspace/book_review');r=base/'Prosistant_50_State_Practice_Books';out=base/'results'
core=json.loads((r/'developer-data/shared_curriculum_1000.json').read_text());states=json.loads((r/'authoring/state_registry.json').read_text());samples=json.loads((base/'samples.json').read_text());logs=json.loads((out/'sources/retrieval_log.json').read_text())
notes={}
for line in (base/'family_notes.tsv').read_text().splitlines():
 n,change,action=line.split('\t');notes[int(n)]=(change,action)
source_map=[['HEADSTART_SUPERVISION','CCDF'],['CCDF'],['CCDF'],['CCDF'],['NIH_SLEEP','CDC_SLEEP','CPSC_CRIB'],['CCDF','HEADSTART'],['CDC_ALLERGY'],['CCDF'],['CDC_DIAPER','CDC_HANDS'],['CDC_CLEAN','EPA'],['CCDF','HEADSTART'],['CCDF','HEADSTART'],['NY_REPORT','CCDF'],['NAEYC_DAP','HEADSTART'],['ADA_CHILDCARE','IDEA','HEADSTART'],['CCDF'],['NHTSA','HEADSTART'],['CCDF'],['CCDF'],['CCDF'],['ADA_CHILDCARE','CCDF'],['CCDF'],['CCDF'],['NAEYC_DAP','CDC_MILESTONES','IDEA'],['CCDF']]
local_tasks=[
'Supervision duties; transition records; missing-child response; exact center/home license rules.',
'Age definitions; staff ratios; group size; own/related-child counting; room and license limits.',
'Background screening; provisional work; role qualifications; volunteer counting; credential recognition.',
'Training approval; delivery mode; renewals; required on-site coverage; orientation.',
'Safe-sleep positioning, surface, equipment, direct observation, records, and any permitted medical exception.',
'Routine authorization; label/instructions; staff training; storage/access; records; emergency medication plans.',
'Allergy plans; cross-contact; food preparation/storage; emergency response; any food-program conditions.',
'Illness exclusion/return criteria; public-health direction; notification triggers and timelines.',
'Hand hygiene; diapering/toileting sequence; permitted surface uses; sanitation and supervision.',
'Label instructions; secure storage; cleaning/sanitizing/disinfecting differences; exposure response.',
'Emergency plans; drills; evacuation/shelter; individual needs; communication and reunification.',
'Care and emergency response; parent notice; licensing notice; reportability; records and deadlines.',
'Reporting trigger; mandated reporter; direct recipient; oral/written timeline; confidentiality; retaliation protections.',
'Prohibited discipline; protective intervention limits; training/reporting; developmentally appropriate support.',
'Civil-rights coverage; individualized modifications; exclusions; Head Start only if applicable; referral links.',
'Attendance corrections; authorized release; current custody documents; late pickup and emergency escalation.',
'Driver/vehicle eligibility; restraints; passenger checks; vehicle-empty checks; staffing; incident reporting.',
'Playground approval; weather; water type; barriers; supervision; rescue readiness; outing permissions.',
'Approved space; fire/building requirements; utilities; hazardous materials; repairs and reopening.',
'License-specific home/household applicability; resident screening; visitors; animals; weapons; notifications.',
'Record-type retention; consent/use limits; access; legal holds; privacy-law coverage.',
'Current subsidy program; authorization dates; enrollment/payment basis; absence payments; family charges; corrections.',
'Ownership/change approval; effective license conditions; inspection authority; corrections and appeals.',
'Rest versus infant sleep; curriculum requirements; family access; child participation; developmental support.',
'Policy effective versions; evidence integrity; source monitoring; corrective follow-up and accountability.']
variant_concerns=[
'Immediate action is not specific enough unless the actual hazard and feasible supervision plan are named.',
'The provider is alone: do not require a second caregiver who is not available. Distinguish tasks that can wait from an emergency.',
'The parent-pressure premise may not fit this incident. Explain the actual non-negotiable safety step and a verified alternative.',
'Name the exact child/task facts that must be handed over; teach-back does not replace eligibility.',
'Do not let reconciling records delay urgent care or an already-required report. Keep observations distinct from assumptions.',
'Distinguish permitted practical help from child-care responsibility, record access, and clearance.',
'Explain what this family is entitled to know while preserving authorized reporting duties. Avoid a universal privacy-law claim.',
'Use a measurable, incident-specific change and test it; do not simply repeat “review the process.”',
'Supply a real primary rule excerpt for a specified category; a research prompt alone cannot provide a scored legal answer.',
'Give event-specific evidence and acknowledge uncertainty; inspection explanations are not a substitute for care or reporting.'
]
high={1,2,3,4,17,18,19,20,21,22,23,24,25,26,33,34,37,38,40,41,45,46,49,50,51,52,53,55,61,63,64,65,66,67,69,71,73,80}
def writecsv(path,rows):
 with path.open('w',newline='',encoding='utf-8-sig') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
families=[];audits=[]
for fi in range(1,101):
 c=core[(fi-1)*10];ti=(fi-1)//4;change,action=notes[fi];priority='P1 safety-critical teaching review' if fi in high else 'P2 legal specificity / instructional revision'
 if fi in (77,78,79,80):change+=' Do not automatically assign the household premise to a commercial center; check applicability first.'
 row=dict(family_id=c['family_id'],core_id_start=c['core_id'],core_id_end=core[fi*10-1]['core_id'],topic=c['topic'],title=c['title'].split(' — ')[0],priority=priority,original_incident=c['scenario'].split('supervision. ')[-1],original_general_action=c['model_answer'].removesuffix(' Address immediate safety and assign eligible coverage before investigating further.'),specific_recommendation=change,provisional_plain_language_action=action,source_leads='; '.join(source_map[ti]),local_verification_needed=local_tasks[ti],source_status='Blocked: primary texts not retrieved',editorial_status='Provisional editorial recommendation; legal and independent bilingual review outstanding')
 families.append(row)
 for vi,c in enumerate(core[(fi-1)*10:fi*10]):
  audits.append(dict(core_id=c['core_id'],family_id=c['family_id'],topic=c['topic'],setting=c['setting'],decision_lens=c['decision_lens'],original_question=c['question'],original_model_answer=c['model_answer'],original_rationale=c['rationale'],question_reuse_count=100,rationale_reuse_count=100,priority=priority,family_recommendation=change,variation_specific_concern=variant_concerns[vi],provisional_action_note=action,proposed_next_step='Rewrite this variant with event-specific facts, question, answer, rationale and rubric; verify applicable primary text before legal scoring.',source_leads='; '.join(source_map[ti]),state_rule_verified=False,review_scope='Template/family-based editorial triage; not individual regulatory approval'))
writecsv(out/'FAMILY_REVISION_PLAN.csv',families);writecsv(out/'CORE_1000_AUDIT.csv',audits)
writecsv(out/'STATE_50_RESEARCH_QUEUE.csv',[dict(state=s['state'],code=s['code'],supplied_entry_point=s['url'],input_reported_source_access=s['source_access'],this_review_access='Blocked by network proxy: 403 Forbidden',mapped_cases=0,exact_license_category='Unresolved',local_authorities='Unresolved',protective_reporting_sources='Unresolved',subsidy_program_sources='Unresolved',next_step='Retrieve issuing-source rules; establish scope/effective date; map propositions and case differences; retain snapshots; obtain qualified review.',regulatory_approval='Not obtained') for s in states])
source_rows=[]
for l in logs:
 k=l['source_id'];scope='State licensing discovery lead; not a case-level citation' if k.startswith('STATE_') else 'Federal, professional or example state source lead; coverage and current text must be verified'
 source_rows.append(dict(source_id=k,requested_url=l['requested_url'],intended_role=scope,retrieved_utc=l['retrieved_utc'],outcome=l.get('error','retrieved'),text_read=False,legal_propositions_verified=0))
writecsv(out/'SOURCE_CATALOG.csv',source_rows)
(out/'SAMPLE_LESSONS_12.json').write_text(json.dumps({'status':'Provisional editorial examples; source research blocked; not an approved exam','audience':'General child-care staff competency','lessons':samples},ensure_ascii=False,indent=2))
css='body{font:17px/1.65 system-ui,sans-serif;max-width:1000px;margin:35px auto;padding:0 22px;color:#193b3a}h1,h2,h3{line-height:1.25}article{border-top:2px solid #c4d7d1;margin-top:32px;padding-top:20px}.notice{background:#fff0d2;padding:18px;border-left:5px solid #a16d12}table{border-collapse:collapse;width:100%;font-size:15px}td,th{border:1px solid #bbb;padding:8px;vertical-align:top}summary{cursor:pointer;font-weight:bold}.meta{font-size:14px;color:#506462}a{color:#075d63} .zh{background:#eef5f3;padding:12px}@media print{body{font-size:11pt;margin:0}article{break-inside:avoid}details{display:block}details>summary{display:none}}'
def page(title,body):return '<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+html.escape(title)+'</title><style>'+css+'</style><body>'+body+'</body></html>'
e=html.escape
notice='<div class="notice"><strong>Provisional editorial draft.</strong> Current authoritative sources could not be accessed. These examples are not verified state-law answers, an approved course, medical instructions, or a validated examination. Respond to real emergencies through the appropriate emergency plan immediately.</div>'
body='<h1>100 foundational teaching prototypes</h1>'+notice+'<p>One prototype per original situation family. This condenses repeated material; it does not replace the pending state/category legal review. Original incident wording is retained. The new action notes are plain-language authoring examples.</p>'
for f in families:
 body+=f'<article id="{f["family_id"]}"><h2>{e(f["family_id"])} · {e(f["title"])}</h2><p class="meta">{e(f["topic"])} · Original range {f["core_id_start"]}–{f["core_id_end"]}</p><h3>Situation</h3><p>{e(f["original_incident"])}</p><h3>Try before reading</h3><p>For this situation, what is the immediate concern and a practical next action? Explain how children remain safely supervised.</p><h3>Provisional teaching note</h3><p>{e(f["provisional_plain_language_action"])}</p><h3>Authoring correction</h3><p>{e(f["specific_recommendation"])}</p><h3>Local facts still needed</h3><p>{e(f["local_verification_needed"])}</p><p class="meta">Unverified source leads: {e(f["source_leads"])}</p></article>'
(out/'FOUNDATION_100_LESSONS.html').write_text(page('100 foundational teaching prototypes',body))
body='<h1>12 worked lessons: study, practice and demonstration</h1>'+notice+'<p>Audience: general child-care staff. Chinese blocks summarize the teaching point; they are not full independently reviewed translations of every field. Incorrect options are teaching distractors, not recommended actions. Source IDs refer to SOURCE_CATALOG.csv; none was verified in this run.</p>'
for i,s in enumerate(samples,1):
 body+=f'<article><h2>{i}. {e(s["title"])}</h2><p class="meta">Original FAM-{s["family"]:03} · new provisional example</p><h3>Situation</h3><p>{e(s["scenario"])}</p><h3>Practice</h3><p>{e(s["question"])}</p><ol type="A">'+''.join('<li>'+e(x)+'</li>' for x in s['options'])+'</ol>'
 body+=f'<details><summary>Show answer and feedback</summary><p><strong>Answer: {s["answer"]}</strong></p><p>{e(s["why"])}</p><p><strong>Everyday example:</strong> {e(s["example"])}</p><p><strong>Demonstrate:</strong> {e(s["task"])}</p><p><strong>Safety feedback:</strong> {e(s["critical"])}</p><p class="zh" lang="zh-Hans">{e(s["zh"])}</p><p class="meta">Sources still to verify: {e(", ".join(s["sources"]))}</p></details></article>'
body+='<script>window.addEventListener("beforeprint",()=>document.querySelectorAll("details").forEach(x=>x.open=true));</script>'
(out/'SAMPLE_LESSONS_12.html').write_text(page('12 provisional worked lessons',body))
# Lightweight Markdown rendering for this controlled report, with real tables in both HTML and Word.
text=(base/'report.md').read_text();(out/'REVIEW_REPORT.md').write_text(text)
doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(11)
sec=doc.sections[0];sec.header.paragraphs[0].text='PROSISTANT · PRELIMINARY EDITORIAL REVIEW · SOURCES PENDING'
blocks=[];lines=text.splitlines();i=0
while i<len(lines):
 line=lines[i]
 if not line.strip():i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   cells=[x.strip() for x in lines[i].strip('|').split('|')]
   if not all(re.fullmatch(r'[-: ]+',x) for x in cells):rows.append(cells)
   i+=1
  table=doc.add_table(rows=0, cols=len(rows[0]));table.style='Light Shading Accent 1'
  for row in rows:
   cells=table.add_row().cells
   for cell,v in zip(cells,row):cell.text=v.replace('**','')
  blocks.append('<table>'+''.join('<tr>'+''.join(('<th>' if n==0 else '<td>')+e(v)+('</th>' if n==0 else '</td>') for v in row)+'</tr>' for n,row in enumerate(rows))+'</table>');continue
 if line.startswith('#'):
  n=len(line)-len(line.lstrip('#'));v=line[n:].strip();doc.add_heading(v,min(n,4));blocks.append(f'<h{n}>{e(v)}</h{n}>')
 else:
  style='List Bullet' if line.startswith('- ') else 'List Number' if re.match(r'^\d+\. ',line) else None
  v=re.sub(r'^(?:- |\d+\. )','',line) if style else line
  p=doc.add_paragraph(style=style)
  for j,part in enumerate(v.split('**')):p.add_run(part).bold=bool(j%2)
  v=e(v);v=re.sub(r'\*\*(.+?)\*\*',r'<strong>\1</strong>',v);blocks.append('<p>'+v+'</p>')
 i+=1
doc.save(out/'REVIEW_REPORT.docx');(out/'REVIEW_REPORT.html').write_text(page('Prosistant preliminary review',''.join(blocks)))
index='<h1>Prosistant review package</h1>'+notice+'<p>Prepared for general child-care staff competency. Start with the report. Current 50-state regulatory verification remains blocked by source access.</p><ul>'
for f,desc in [('REVIEW_REPORT.html','Read the review and priorities'),('REVIEW_REPORT.docx','Download the Word report'),('FAMILY_REVISION_PLAN.csv','100 family revision recommendations'),('CORE_1000_AUDIT.csv','1,000 shared-exercise audit rows'),('FOUNDATION_100_LESSONS.html','100 condensed teaching prototypes'),('SAMPLE_LESSONS_12.html','12 worked teaching examples'),('STATE_50_RESEARCH_QUEUE.csv','50-state legal research queue'),('SOURCE_CATALOG.csv','Source access and verification status'),('structural_audit.json','Structural audit evidence')]:index+=f'<li><a href="{f}">{e(desc)}</a></li>'
index+='</ul><p>Original materials were preserved. No approval, credit, publication, or validated examination status is claimed.</p>'
(out/'START_HERE.html').write_text(page('Prosistant review package',index))
# Retain reproducible audit/research helpers, without the original 190 MB extracted archive.
for name in ['audit.py','research.py','package_review.py']:
 (out/name).write_text((base/name).read_text())
manifest={}
for p in sorted(out.rglob('*')):
 if p.is_file() and p.name!='deliverable_manifest.json':manifest[str(p.relative_to(out))]={'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
(out/'deliverable_manifest.json').write_text(json.dumps(manifest,indent=2))
zpath=base/'Prosistant_Preliminary_Review_and_Revision.zip'
with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
 for p in sorted(out.rglob('*')):
  if p.is_file():z.write(p,'Prosistant_Review/'+str(p.relative_to(out)))
print('Created',zpath,'bytes',zpath.stat().st_size)
print('Family recommendations',len(families),'Core audit rows',len(audits),'Sample lessons',len(samples),'Source requests',len(logs))
